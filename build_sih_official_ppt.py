import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def update_presentation():
    input_path = r"D:\sih\SIH26191 BY NOVACORE.pptx.backup"
    prs = Presentation(input_path)

    # Color Palette
    DARK_BLUE = RGBColor(10, 25, 47)      # #0A192F
    NAVY_CARD = RGBColor(14, 37, 55)      # #0E2537
    LIGHT_BG = RGBColor(248, 250, 252)    # #F8FAFC
    CARD_BG = RGBColor(241, 245, 249)     # #F1F5F9
    BORDER_COL = RGBColor(203, 213, 225)  # #CBD5E1
    TEXT_DARK = RGBColor(15, 23, 42)      # #0F172A
    TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
    PRIMARY_CYAN = RGBColor(0, 130, 200)  # #0082C8
    ACCENT_RED = RGBColor(229, 72, 77)    # #E5484D
    ACCENT_GREEN = RGBColor(16, 185, 129) # #10B981
    ACCENT_ORANGE = RGBColor(245, 158, 11)# #F59E0B
    WHITE = RGBColor(255, 255, 255)
    PLACEHOLDER_FILL = RGBColor(235, 243, 252)
    PLACEHOLDER_BORDER = RGBColor(37, 99, 235)

    def remove_shape(shape):
        sp = shape._element
        sp.getparent().remove(sp)

    def add_placeholder_box(slide, left, top, width, height, title, instruction):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = PLACEHOLDER_FILL
        box.line.color.rgb = PLACEHOLDER_BORDER
        box.line.width = Pt(1.5)
        
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.15)
        tf.margin_bottom = Inches(0.15)
        
        p = tf.paragraphs[0]
        p.text = "📷 [PLACE PHOTO HERE]"
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Calibri"
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_CYAN
        
        p2 = tf.add_paragraph()
        p2.text = title
        p2.alignment = PP_ALIGN.CENTER
        p2.font.name = "Calibri"
        p2.font.size = Pt(10.5)
        p2.font.bold = True
        p2.font.color.rgb = DARK_BLUE
        
        p3 = tf.add_paragraph()
        p3.text = f"👉 {instruction}"
        p3.alignment = PP_ALIGN.CENTER
        p3.font.name = "Calibri"
        p3.font.size = Pt(8.5)
        p3.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (UPDATE DETAILS)
    # =========================================================================
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.name == "TextBox 9" and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            
            bullets = [
                ("• Problem Statement ID – ", "SIH26191"),
                ("• Problem Statement Title – ", "Intelligent Identification of Hazard-Based Red Zones, Carrying Capacity Assessment, and Immediate Relocation Needs for Vulnerable Habitations."),
                ("• Theme – ", "Disaster Management"),
                ("• PS Category – ", "Software"),
                ("• Team ID – ", "170297"),
                ("• Team Name – ", "NOVACORE"),
            ]
            for i, (b_title, b_val) in enumerate(bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(10)
                r1 = p.add_run()
                r1.text = b_title
                r1.font.name = "Calibri"
                r1.font.size = Pt(16)
                r1.font.bold = True
                r1.font.color.rgb = TEXT_DARK
                
                r2 = p.add_run()
                r2.text = b_val
                r2.font.name = "Calibri"
                r2.font.size = Pt(16)
                r2.font.bold = (i in [0, 1, 4, 5])
                r2.font.color.rgb = PRIMARY_CYAN if i in [0, 4, 5] else TEXT_DARK

    # =========================================================================
    # SLIDE 2: PROPOSED SOLUTION
    # =========================================================================
    s2 = prs.slides[1]
    # Update title
    for shape in list(s2.shapes):
        if shape.name == "Title 1" and shape.has_text_frame:
            shape.text_frame.text = "PROPOSED SOLUTION: RAKSHASETU (Smart Disaster Intelligence)"
            shape.text_frame.paragraphs[0].font.name = "Calibri"
            shape.text_frame.paragraphs[0].font.size = Pt(20)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name in ["TextBox 15361", "TextBox 15362", "Picture 3"]:
            remove_shape(shape)

    # Subtitle Link Bar across top
    sub2 = s2.shapes.add_textbox(Inches(0.6), Inches(1.15), Inches(12.13), Inches(0.4))
    stf2 = sub2.text_frame
    stf2.margin_left = stf2.margin_top = stf2.margin_right = stf2.margin_bottom = 0
    sp2 = stf2.paragraphs[0]
    sp2.text = "RAKSHASETU — AI-Powered Hazard-Zone, Vulnerability & Relocation Intelligence Platform  |  GitHub: https://github.com/vinitmishraaa/RAKSHASETU"
    sp2.font.name = "Calibri"
    sp2.font.size = Pt(10.5)
    sp2.font.bold = True
    sp2.font.color.rgb = PRIMARY_CYAN

    # Left Column: 3 Pillar Cards (matching KAWA PDF 1 style)
    c_w = Inches(5.8)
    l_x = Inches(0.6)

    # Card 1: Red Zones
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_x, Inches(1.6), c_w, Inches(1.65))
    c1.fill.solid()
    c1.fill.fore_color.rgb = WHITE
    c1.line.color.rgb = ACCENT_RED
    c1.line.width = Pt(1.5)
    ctf1 = c1.text_frame
    ctf1.word_wrap = True
    ctf1.margin_left = ctf1.margin_right = ctf1.margin_top = ctf1.margin_bottom = Inches(0.1)
    cp1 = ctf1.paragraphs[0]
    cp1.text = "🔴 1. Dynamic Hazard-Based Red-Zone Delineation"
    cp1.font.name = "Calibri"
    cp1.font.size = Pt(12)
    cp1.font.bold = True
    cp1.font.color.rgb = ACCENT_RED
    for text in [
        "Identifies zones Unsuitable for Permanent Habitation dynamically via live sensor feeds.",
        "4 PS Hazards Modeled: Floods, Landslides, Coastal Erosion, and Cloudbursts.",
        "GIS Animated Blowout Rings: Pulsating radar perimeters pinpoint threat severity.",
    ]:
        bp = ctf1.add_paragraph()
        bp.text = f" {text}"
        bp.font.name = "Calibri"
        bp.font.size = Pt(9.0)
        bp.font.color.rgb = TEXT_DARK

    # Card 2: 3-Tier Relocation
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_x, Inches(3.35), c_w, Inches(1.65))
    c2.fill.solid()
    c2.fill.fore_color.rgb = WHITE
    c2.line.color.rgb = PRIMARY_CYAN
    c2.line.width = Pt(1.5)
    ctf2 = c2.text_frame
    ctf2.word_wrap = True
    ctf2.margin_left = ctf2.margin_right = ctf2.margin_top = ctf2.margin_bottom = Inches(0.1)
    cp2 = ctf2.paragraphs[0]
    cp2.text = "⏱️ 2. 3-Tier Relocation Need Prioritization"
    cp2.font.name = "Calibri"
    cp2.font.size = Pt(12)
    cp2.font.bold = True
    cp2.font.color.rgb = PRIMARY_CYAN
    for text in [
        "🔴 Immediate (0–48 Hours): Rapid tactical evacuation for acute life-threats.",
        "🟠 Short-Term (1–3 Months): Pre-monsoon planned relocation before floods/slides.",
        "🟡 Medium-Term (6–12 Months): Sustainable rehabilitation & permanent housing.",
    ]:
        bp = ctf2.add_paragraph()
        bp.text = f" {text}"
        bp.font.name = "Calibri"
        bp.font.size = Pt(9.0)
        bp.font.color.rgb = TEXT_DARK

    # Card 3: Carrying Capacity & Zero Cost
    c3 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_x, Inches(5.1), c_w, Inches(1.65))
    c3.fill.solid()
    c3.fill.fore_color.rgb = WHITE
    c3.line.color.rgb = ACCENT_GREEN
    c3.line.width = Pt(1.5)
    ctf3 = c3.text_frame
    ctf3.word_wrap = True
    ctf3.margin_left = ctf3.margin_right = ctf3.margin_top = ctf3.margin_bottom = Inches(0.1)
    cp3 = ctf3.paragraphs[0]
    cp3.text = "🏢 3. Carrying Capacity Assessment & Overcrowding Mitigation"
    cp3.font.name = "Calibri"
    cp3.font.size = Pt(12)
    cp3.font.bold = True
    cp3.font.color.rgb = ACCENT_GREEN
    for text in [
        "Evaluates Total Capacity, Pre-Occupancy, and Net Absorption Headroom.",
        "Automated Multi-Site Split: Prevents secondary disaster outbreak when stress >85%.",
        "100% Zero Operating Cost: Built completely on open-source APIs (Open-Meteo, USGS, OSM).",
    ]:
        bp = ctf3.add_paragraph()
        bp.text = f" {text}"
        bp.font.name = "Calibri"
        bp.font.size = Pt(9.0)
        bp.font.color.rgb = TEXT_DARK

    # Right Column: Links Box + 2 Photo Placeholders
    r_x = Inches(6.7)
    r_w = Inches(6.0)

    # Project Links Box
    lbox = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, r_x, Inches(1.6), r_w, Inches(0.85))
    lbox.fill.solid()
    lbox.fill.fore_color.rgb = NAVY_CARD
    lbox.line.color.rgb = DARK_BLUE
    ltf = lbox.text_frame
    ltf.margin_top = Inches(0.08)
    lp1 = ltf.paragraphs[0]
    lp1.text = "📱 PROJECT LINKS & REPOSITORIES (Click to Open):"
    lp1.font.name = "Calibri"
    lp1.font.size = Pt(10.5)
    lp1.font.bold = True
    lp1.font.color.rgb = PRIMARY_CYAN
    lp2 = ltf.add_paragraph()
    lp2.text = "• GitHub Code Repository: https://github.com/vinitmishraaa/RAKSHASETU"
    lp2.font.name = "Calibri"
    lp2.font.size = Pt(9.0)
    lp2.font.color.rgb = WHITE
    lp3 = ltf.add_paragraph()
    lp3.text = "• Live Web Dashboard: http://localhost:5173  |  FastAPI Docs: /docs"
    lp3.font.name = "Calibri"
    lp3.font.size = Pt(9.0)
    lp3.font.color.rgb = WHITE

    add_placeholder_box(
        s2, r_x, Inches(2.55), r_w, Inches(2.15),
        "Main Dashboard: Multi-Hazard Red-Zone GIS Map",
        "Screenshot from http://localhost:5173 with State='West Bengal' and Red Zones filter ON."
    )
    add_placeholder_box(
        s2, r_x, Inches(4.8), r_w, Inches(1.95),
        "Carrying Capacity Assessment & Stress Gauge Card",
        "Screenshot from Relocation page showing capacity progress bar, stress %, and headroom."
    )

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH
    # =========================================================================
    s3 = prs.slides[2]
    for shape in list(s3.shapes):
        if shape.name == "Title 1" and shape.has_text_frame:
            shape.text_frame.text = "TECHNICAL APPROACH: End-to-End System Workflow"
            shape.text_frame.paragraphs[0].font.name = "Calibri"
            shape.text_frame.paragraphs[0].font.size = Pt(20)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name in ["TextBox 17409", "TextBox 17410", "TextBox 17411", "Picture 3", "AutoShape 2"]:
            remove_shape(shape)

    # Top 5 Steps Flow Boxes
    steps = [
        ("Step 1: Telemetry", "Live Open Data\n• Open-Meteo Weather\n• USGS Earthquakes\n• NDMA SACHET CAP"),
        ("Step 2: Risk Engine", "3-Pillar Formula\n• Hazard Intensity (40%)\n• Vulnerability (35%)\n• Disaster History (25%)"),
        ("Step 3: Red Zones", "Red-Zone Delineation\n• Multi-hazard unsuitability\n• 4 Hazards: Flood, Slide,\n  Erosion, Cloudburst"),
        ("Step 4: Safe Sites", "Capacity Assessment\n• Headroom & stress %\n• Overcrowding split\n• OSRM road routing"),
        ("Step 5: Action", "SDMA Action Corridor\n• 4-Step guided flow\n• Google Maps driving\n• 🖨️ Printable order"),
    ]
    st_w = Inches(2.3)
    st_gap = Inches(0.15)
    for i, (title, desc) in enumerate(steps):
        sx = Inches(0.6) + i * (st_w + st_gap)
        sb = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, Inches(1.3), st_w, Inches(1.4))
        sb.fill.solid()
        sb.fill.fore_color.rgb = WHITE
        sb.line.color.rgb = PRIMARY_CYAN if i % 2 == 0 else ACCENT_RED
        sb.line.width = Pt(1.5)
        sbtf = sb.text_frame
        sbtf.margin_left = sbtf.margin_right = sbtf.margin_top = sbtf.margin_bottom = Inches(0.08)
        p1 = sbtf.paragraphs[0]
        p1.text = title
        p1.font.name = "Calibri"
        p1.font.size = Pt(10.5)
        p1.font.bold = True
        p1.font.color.rgb = DARK_BLUE
        p2 = sbtf.add_paragraph()
        p2.text = desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(8.0)
        p2.font.color.rgb = TEXT_DARK

    # Bottom Left: Tech Stack
    tcard = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(2.85), Inches(5.8), Inches(3.95))
    tcard.fill.solid()
    tcard.fill.fore_color.rgb = WHITE
    tcard.line.color.rgb = BORDER_COL
    tctf = tcard.text_frame
    tctf.margin_left = tctf.margin_right = tctf.margin_top = tctf.margin_bottom = Inches(0.12)
    tp1 = tctf.paragraphs[0]
    tp1.text = "💻 TECHNOLOGY STACK & SMART FEATURES:"
    tp1.font.name = "Calibri"
    tp1.font.size = Pt(12)
    tp1.font.bold = True
    tp1.font.color.rgb = DARK_BLUE
    for tb in [
        "Frontend: React 18, TypeScript, Vite, Leaflet.js, Recharts, Custom Glassmorphism UI.",
        "Backend: Python 3.12, FastAPI, In-Memory Administrative Directory (36 States/UTs).",
        "Road Network Routing: Open Source Routing Machine (OSRM) driving route geometry & live ETA.",
        "3-Pillar Risk Engine: 40% Hazard Intensity + 35% Vulnerability + 25% Disaster History Recurrence.",
        "Carrying Capacity Engine: Dynamic absorption headroom gauge & automated split allocation.",
        "SDMA Decision Suite: 4-Step evacuation corridor, Google Maps mobile handoff, printable order brief.",
        "GitHub Code Link: https://github.com/vinitmishraaa/RAKSHASETU",
    ]:
        tbp = tctf.add_paragraph()
        tbp.text = f"• {tb}"
        tbp.font.name = "Calibri"
        tbp.font.size = Pt(9.0)
        tbp.font.color.rgb = TEXT_DARK

    # Bottom Right: 2 Screenshot Placeholders
    add_placeholder_box(
        s3, Inches(6.7), Inches(2.85), Inches(6.0), Inches(1.9),
        "4-Step Guided Official Relocation Corridor & Road Route",
        "Screenshot from Relocation page showing 4 steps, OSRM road distance, ETA, and turn-by-turn guidance."
    )
    add_placeholder_box(
        s3, Inches(6.7), Inches(4.85), Inches(6.0), Inches(1.95),
        "Official SDMA Relocation Order Print Preview",
        "Screenshot of the print dialog after clicking 'Export / Print Official SDMA Relocation Order'."
    )

    # =========================================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # =========================================================================
    s4 = prs.slides[3]
    for shape in list(s4.shapes):
        if shape.name == "Title 1" and shape.has_text_frame:
            shape.text_frame.text = "FEASIBILITY & VIABILITY: Operational & Comparative Assessment"
            shape.text_frame.paragraphs[0].font.name = "Calibri"
            shape.text_frame.paragraphs[0].font.size = Pt(20)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name in ["Picture 4"]:
            remove_shape(shape)

    # Table on Left
    table_shape = s4.shapes.add_table(6, 4, Inches(0.6), Inches(1.3), Inches(7.0), Inches(3.4))
    table = table_shape.table
    table.columns[0].width = Inches(1.5)
    table.columns[1].width = Inches(1.9)
    table.columns[2].width = Inches(2.1)
    table.columns[3].width = Inches(1.5)

    table_data = [
        ["Operational Dimension", "Traditional Reactive Method", "RakshaSetu Proactive AI-GIS", "Key Advantage"],
        ["Response Timing", "Post-disaster rescue (Panic & delay)", "Proactive pre-disaster tiers (0-48h, 1-3m, 6-12m)", "Zero panic & loss"],
        ["Hazard Delineation", "Manual, localized post-event survey", "Dynamic GIS Multi-Hazard Red-Zone mapping", "High precision"],
        ["Safe Site Allocation", "Ad-hoc intake (Severe overcrowding)", "Carrying capacity headroom stress % analysis", "Prevents stampede"],
        ["Evacuation Routing", "Unverified paths (often blocked)", "OSRM real road network & Google Maps handoff", "Guaranteed ETA"],
        ["Financial Cost", "Expensive proprietary GIS licenses", "100% Open-source zero-key architecture", "₹0 Software Bills"],
    ]
    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            cp = cell.text_frame.paragraphs[0]
            cp.font.name = "Calibri"
            cp.font.size = Pt(8.5 if r_idx > 0 else 9.5)
            cp.font.bold = (r_idx == 0 or c_idx == 0)
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = DARK_BLUE
                cp.font.color.rgb = WHITE
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r_idx % 2 == 0 else CARD_BG
                cp.font.color.rgb = TEXT_DARK

    # Bottom Left: 3 Reasons Box
    rbox = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(4.85), Inches(7.0), Inches(1.95))
    rbox.fill.solid()
    rbox.fill.fore_color.rgb = WHITE
    rbox.line.color.rgb = ACCENT_ORANGE
    rbox.line.width = Pt(1.5)
    rtf = rbox.text_frame
    rtf.margin_left = rtf.margin_right = rtf.margin_top = rtf.margin_bottom = Inches(0.1)
    rp = rtf.paragraphs[0]
    rp.text = "🎯 3 REASONS WHY RAKSHASETU IS 100% FEASIBLE & VIABLE:"
    rp.font.name = "Calibri"
    rp.font.size = Pt(11)
    rp.font.bold = True
    rp.font.color.rgb = ACCENT_ORANGE
    for r in [
        "Zero Key Operating Cost: Relies on public Open-Meteo, USGS, NDMA CAP and OpenStreetMap with no cloud bills.",
        "Deterministic & Explainable AI: Formulated on clear mathematical equations (40% Hazard, 35% Vulnerability, 25% History) avoiding black-box errors.",
        "Pan-India Immediate Deployment: All 28 States and 8 UTs pre-indexed; instantly scalable to every SDMA and DDMA.",
    ]:
        rr = rtf.add_paragraph()
        rr.text = f" {r}"
        rr.font.name = "Calibri"
        rr.font.size = Pt(8.5)
        rr.font.color.rgb = TEXT_DARK

    # Right: Photo Placeholders
    add_placeholder_box(
        s4, Inches(7.8), Inches(1.3), Inches(4.9), Inches(2.65),
        "Safe Sites Infrastructure & Headroom Stress Meter",
        "Screenshot from Safe Sites page (/safe-sites) showing shelter cards with capacity, occupancy, and headroom."
    )
    add_placeholder_box(
        s4, Inches(7.8), Inches(4.05), Inches(4.9), Inches(2.75),
        "Multi-Site Split Allocation to Prevent Overcrowding",
        "Screenshot showing Relocation page recommendation when high population is allocated safely across shelters."
    )

    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS
    # =========================================================================
    s5 = prs.slides[4]
    for shape in list(s5.shapes):
        if shape.name == "Title 1" and shape.has_text_frame:
            shape.text_frame.text = "IMPACT & BENEFITS: Saving Lives, Preventing Overcrowding & Empowering SDMA"
            shape.text_frame.paragraphs[0].font.name = "Calibri"
            shape.text_frame.paragraphs[0].font.size = Pt(20)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name in ["TextBox 17409", "TextBox 17410", "Picture 2"]:
            remove_shape(shape)

    # Left Column: 3 Impact Cards + Measurable Outcomes
    im_w = Inches(5.8)
    
    im1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.3), im_w, Inches(1.2))
    im1.fill.solid()
    im1.fill.fore_color.rgb = WHITE
    im1.line.color.rgb = ACCENT_RED
    itf1 = im1.text_frame
    itf1.margin_left = itf1.margin_right = itf1.margin_top = itf1.margin_bottom = Inches(0.08)
    ip1 = itf1.paragraphs[0]
    ip1.text = "🛡️ 1. Life & Infrastructure Protection (Shifting Reactive to Proactive)"
    ip1.font.name = "Calibri"
    ip1.font.size = Pt(11)
    ip1.font.bold = True
    ip1.font.color.rgb = ACCENT_RED
    ib1 = itf1.add_paragraph()
    ib1.text = "Eliminates repeated loss of life and property in perennial flood, landslide, erosion, and cloudburst belts by moving vulnerable habitations BEFORE disaster strikes."
    ib1.font.name = "Calibri"
    ib1.font.size = Pt(8.5)
    ib1.font.color.rgb = TEXT_DARK

    im2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(2.6), im_w, Inches(1.2))
    im2.fill.solid()
    im2.fill.fore_color.rgb = WHITE
    im2.line.color.rgb = PRIMARY_CYAN
    itf2 = im2.text_frame
    itf2.margin_left = itf2.margin_right = itf2.margin_top = itf2.margin_bottom = Inches(0.08)
    ip2 = itf2.paragraphs[0]
    ip2.text = "🏢 2. Prevention of Secondary Disasters & Post-Evacuation Overcrowding"
    ip2.font.name = "Calibri"
    ip2.font.size = Pt(11)
    ip2.font.bold = True
    ip2.font.color.rgb = PRIMARY_CYAN
    ib2 = itf2.add_paragraph()
    ib2.text = "Mathematical carrying capacity assessment caps intake at safe thresholds, preventing sanitation collapse, epidemic outbreaks, and resource exhaustion at relief centres."
    ib2.font.name = "Calibri"
    ib2.font.size = Pt(8.5)
    ib2.font.color.rgb = TEXT_DARK

    im3 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(3.9), im_w, Inches(1.2))
    im3.fill.solid()
    im3.fill.fore_color.rgb = WHITE
    im3.line.color.rgb = ACCENT_GREEN
    itf3 = im3.text_frame
    itf3.margin_left = itf3.margin_right = itf3.margin_top = itf3.margin_bottom = Inches(0.08)
    ip3 = itf3.paragraphs[0]
    ip3.text = "📋 3. Empowering SDMA, DDMA & Incident Commanders with Action Orders"
    ip3.font.name = "Calibri"
    ip3.font.size = Pt(11)
    ip3.font.bold = True
    ip3.font.color.rgb = ACCENT_GREEN
    ib3 = itf3.add_paragraph()
    ib3.text = "Generates signed official relocation briefs, driver turn-by-turn routes, and automated SMS/Email alert dispatch to pre-configured disaster response officers."
    ib3.font.name = "Calibri"
    ib3.font.size = Pt(8.5)
    ib3.font.color.rgb = TEXT_DARK

    # Outcomes Box
    pbox = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.2), im_w, Inches(1.6))
    pbox.fill.solid()
    pbox.fill.fore_color.rgb = NAVY_CARD
    pbox.line.color.rgb = DARK_BLUE
    ptf = pbox.text_frame
    ptf.margin_left = ptf.margin_right = ptf.margin_top = ptf.margin_bottom = Inches(0.08)
    pp = ptf.paragraphs[0]
    pp.text = "📊 PROVED MEASURABLE OUTCOMES (Per Vulnerable District):"
    pp.font.name = "Calibri"
    pp.font.size = Pt(10.5)
    pp.font.bold = True
    pp.font.color.rgb = PRIMARY_CYAN
    for o in [
        "0–48 Hours Evacuation Window: Immediate alerts for high-risk red-zone communities.",
        "85% Carrying Capacity Threshold: Triggers multi-site diversion to protect shelter safety.",
        "100% Open Data Transparency: Zero recurring software or mapping license burden.",
        "36 States & UTs Pre-Indexed: Ready for instant pan-India adoption by NDRF / MHA.",
    ]:
        op = ptf.add_paragraph()
        op.text = f"• {o}"
        op.font.name = "Calibri"
        op.font.size = Pt(8.0)
        op.font.color.rgb = WHITE

    # Right: 2 Screenshot Placeholders
    add_placeholder_box(
        s5, Inches(6.7), Inches(1.3), Inches(6.0), Inches(2.65),
        "Recharts Live Analytics: 3-Tier Distribution & 4-Hazard Profile",
        "Screenshot from Analytics page (/analytics) showing 3-Tier Pie Chart and 4-Hazard Bar Chart."
    )
    add_placeholder_box(
        s5, Inches(6.7), Inches(4.1), Inches(6.0), Inches(2.7),
        "Officer Alert Desk & Dual-Tone Emergency Siren Tester",
        "Screenshot from Alerts page (/alerts) showing 5 response officers with Email/SMS dispatch and siren tester."
    )

    # =========================================================================
    # SLIDE 6: RESEARCH, PROJECT LINKS & FUTURE ROADMAP
    # =========================================================================
    s6 = prs.slides[5]
    for shape in list(s6.shapes):
        if shape.name == "Title 1" and shape.has_text_frame:
            shape.text_frame.text = "RESEARCH, PROJECT LINKS & FUTURE ROADMAP"
            shape.text_frame.paragraphs[0].font.name = "Calibri"
            shape.text_frame.paragraphs[0].font.size = Pt(20)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.name in ["TextBox 17409", "TextBox 17410", "TextBox 17411", "Picture 2"]:
            remove_shape(shape)

    # Top Links Box
    top_links = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.3), Inches(12.13), Inches(0.85))
    top_links.fill.solid()
    top_links.fill.fore_color.rgb = WHITE
    top_links.line.color.rgb = PRIMARY_CYAN
    top_links.line.width = Pt(1.5)
    tltf = top_links.text_frame
    tltf.margin_top = Inches(0.06)
    tlp = tltf.paragraphs[0]
    tlp.text = "📱 OFFICIAL PROJECT DELIVERABLES & REPOSITORIES:"
    tlp.font.name = "Calibri"
    tlp.font.size = Pt(11)
    tlp.font.bold = True
    tlp.font.color.rgb = PRIMARY_CYAN
    tlp2 = tltf.add_paragraph()
    tlp2.text = "• GitHub Code Repository: https://github.com/vinitmishraaa/RAKSHASETU  |  Branch: main (Tested & Verified)"
    tlp2.font.name = "Calibri"
    tlp2.font.size = Pt(9.5)
    tlp2.font.color.rgb = TEXT_DARK
    tlp3 = tltf.add_paragraph()
    tlp3.text = "• Local Command Center: http://localhost:5173  |  FastAPI Backend REST Endpoints: http://127.0.0.1:8000/docs"
    tlp3.font.name = "Calibri"
    tlp3.font.size = Pt(9.5)
    tlp3.font.color.rgb = TEXT_DARK

    # 3 Columns for Sources
    col_w = Inches(3.9)
    col_gap = Inches(0.2)
    
    # Col 1: Government Sources
    c1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(2.3), col_w, Inches(3.3))
    c1.fill.solid()
    c1.fill.fore_color.rgb = WHITE
    c1.line.color.rgb = DARK_BLUE
    c1tf = c1.text_frame
    c1tf.margin_left = c1tf.margin_right = c1tf.margin_top = c1tf.margin_bottom = Inches(0.1)
    cp1 = c1tf.paragraphs[0]
    cp1.text = "🏛️ Official Government Sources"
    cp1.font.name = "Calibri"
    cp1.font.size = Pt(11)
    cp1.font.bold = True
    cp1.font.color.rgb = DARK_BLUE
    for name, url in [
        ("Ministry of Home Affairs (MHA)", "https://mha.gov.in"),
        ("National Disaster Response Force (NDRF)", "https://ndrf.gov.in"),
        ("NDMA SACHET (CAP Early Warning)", "https://sachet.ndma.gov.in"),
        ("India Meteorological Department (IMD)", "https://mausam.imd.gov.in"),
        ("Geological Survey of India (GSI - Landslides)", "https://gsi.gov.in"),
        ("Central Water Commission (CWC - Floods)", "https://cwc.gov.in"),
        ("Census of India (Village Demographics)", "https://censusindia.gov.in"),
    ]:
        p = c1tf.add_paragraph()
        p.text = f"• {name}\n  {url}"
        p.font.name = "Calibri"
        p.font.size = Pt(7.5)
        p.font.color.rgb = TEXT_MUTED

    # Col 2: Open / Technical Sources
    c2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6) + col_w + col_gap, Inches(2.3), col_w, Inches(3.3))
    c2.fill.solid()
    c2.fill.fore_color.rgb = WHITE
    c2.line.color.rgb = PRIMARY_CYAN
    c2tf = c2.text_frame
    c2tf.margin_left = c2tf.margin_right = c2tf.margin_top = c2tf.margin_bottom = Inches(0.1)
    cp2 = c2tf.paragraphs[0]
    cp2.text = "🌐 Open Technical Data Sources"
    cp2.font.name = "Calibri"
    cp2.font.size = Pt(11)
    cp2.font.bold = True
    cp2.font.color.rgb = PRIMARY_CYAN
    for name, url in [
        ("Open-Meteo API (Live Weather & Wind)", "https://open-meteo.com"),
        ("USGS Seismology (Earthquake Hypocenters)", "https://earthquake.usgs.gov"),
        ("OpenStreetMap (Roads & Critical Buildings)", "https://openstreetmap.org"),
        ("OSRM Project (Road Routing Engine)", "http://project-osrm.org"),
        ("GDACS (Global Disaster Alert & Coordination)", "https://gdacs.org"),
        ("NASA GIBS (Satellite True-Color Imagery)", "https://earthdata.nasa.gov"),
        ("Leaflet.js & Recharts (Visualization)", "https://leafletjs.com"),
    ]:
        p = c2tf.add_paragraph()
        p.text = f"• {name}\n  {url}"
        p.font.name = "Calibri"
        p.font.size = Pt(7.5)
        p.font.color.rgb = TEXT_MUTED

    # Col 3: Research & Development Sources
    c3 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6) + 2 * (col_w + col_gap), Inches(2.3), col_w, Inches(3.3))
    c3.fill.solid()
    c3.fill.fore_color.rgb = WHITE
    c3.line.color.rgb = ACCENT_RED
    c3tf = c3.text_frame
    c3tf.margin_left = c3tf.margin_right = c3tf.margin_top = c3tf.margin_bottom = Inches(0.1)
    cp3 = c3tf.paragraphs[0]
    cp3.text = "📚 Research & Standards"
    cp3.font.name = "Calibri"
    cp3.font.size = Pt(11)
    cp3.font.bold = True
    cp3.font.color.rgb = ACCENT_RED
    for name, desc in [
        ("UNDRR Sendai Framework 2015–2030", "Target 4: Disaster Resilient Infrastructure"),
        ("IPCC AR6 Vulnerability Guidelines", "Climate-induced resettlement standards"),
        ("Smart India Hackathon 2026 Portal", "Problem Statement ID: SIH26191"),
        ("NDMA SOP on Disaster Relocation", "Standard capacity allocations per person"),
        ("Team NOVACORE Research Paper", "Proactive GIS Carrying Capacity Models"),
        ("Open Source MIT License", "Fully libre software for public good"),
    ]:
        p = c3tf.add_paragraph()
        p.text = f"• {name}\n  {desc}"
        p.font.name = "Calibri"
        p.font.size = Pt(7.5)
        p.font.color.rgb = TEXT_MUTED

    # Bottom Banner
    bot_banner = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.75), Inches(12.13), Inches(1.05))
    bot_banner.fill.solid()
    bot_banner.fill.fore_color.rgb = CARD_BG
    bot_banner.line.color.rgb = BORDER_COL
    btf = bot_banner.text_frame
    btf.margin_top = Inches(0.06)
    bp = btf.paragraphs[0]
    bp.text = "🌟 HOW THESE SCIENTIFIC & TECHNICAL REFERENCES VALIDATE RAKSHASETU:"
    bp.font.name = "Calibri"
    bp.font.size = Pt(10)
    bp.font.bold = True
    bp.font.color.rgb = DARK_BLUE
    bp2 = btf.add_paragraph()
    bp2.text = "✅ Reliable Open Data: Official government and live sensor feeds ensure real-time hazard detection without manual delay.\n✅ Proven Mathematical Methodology: 3-Pillar risk formula and carrying capacity headroom calculations ensure zero subjective bias.\n✅ Immediate Administrative Adoption: Structured to fit directly into SDMA / DDMA standard operating procedures for MHA / NDRF."
    bp2.font.name = "Calibri"
    bp2.font.size = Pt(8.0)
    bp2.font.color.rgb = TEXT_DARK

    # Save directly to D:\sih\SIH26191 BY NOVACORE.pptx and also in rakshasetu project folder!
    output_target1 = r"D:\sih\SIH26191 BY NOVACORE.pptx"
    output_target2 = r"D:\sih\rakshasetu\rakshasetu\RAKSHASETU_SIH26191_Official_Template.pptx"
    
    prs.save(output_target1)
    print(f"SUCCESS: Saved main SIH template file at: {output_target1}")
    prs.save(output_target2)
    print(f"SUCCESS: Saved project copy at: {output_target2}")

    try:
        locked_target = r"D:\sih\rakshasetu\rakshasetu\RAKSHASETU_SIH26191_Presentation.pptx"
        prs.save(locked_target)
        print(f"SUCCESS: Also updated: {locked_target}")
    except Exception as e:
        print(f"Note: {locked_target} is currently open in PowerPoint ({e}).")

if __name__ == "__main__":
    update_presentation()
